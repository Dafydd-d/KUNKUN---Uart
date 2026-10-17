"""Excel 导出核心：与界面框架无关的纯写盘逻辑，Tk/Qt 版共用。"""
import os
import sys


def _find_lib_base():
    """自动查找 _lib 依赖的父目录，兼容所有打包模式"""
    lib_name = '_lib'

    # 1. PyInstaller 标准路径：onefile 临时目录 / 新版 onedir 的 _internal 目录
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        meipass_lib = os.path.join(sys._MEIPASS, lib_name)
        if os.path.exists(meipass_lib):
            return sys._MEIPASS

    # 2. 兼容目录模式：exe 同级目录下的 _internal/_lib
    exe_dir = os.path.dirname(sys.executable)
    internal_lib = os.path.join(exe_dir, '_internal', lib_name)
    if os.path.exists(internal_lib):
        return os.path.join(exe_dir, '_internal')

    # 3. 兼容绿色分发：exe 同级直接放 _lib 目录
    exe_lib = os.path.join(exe_dir, lib_name)
    if os.path.exists(exe_lib):
        return exe_dir

    # 4. 源码运行模式：脚本所在目录
    return os.path.dirname(os.path.abspath(__file__))


def ensure_lib_path():
    """将内置 _lib 依赖库加入 Python 导入路径（openpyxl 等）"""
    base_path = _find_lib_base()
    lib_path = os.path.join(base_path, '_lib')
    if os.path.exists(lib_path) and lib_path not in sys.path:
        sys.path.insert(0, lib_path)


def export_frames(data_frames, file_path, labels):
    """把收发数据帧写成 Excel：时间戳 + 方向 + 每字节一列。

    labels 需包含: sheet_name / col_timestamp / col_direction / byte_label。
    openpyxl 缺失时抛 ImportError，写盘失败抛其它异常，由调用方提示。
    """
    ensure_lib_path()
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = labels["sheet_name"]

    # 写入表头：第1列时间戳，第2列方向，第3列起是字节
    max_byte_len = max(len(frame[2]) for frame in data_frames)
    header = [labels["col_timestamp"], labels["col_direction"]]
    header.extend(f"{labels['byte_label']}{i:02d}" for i in range(1, max_byte_len + 1))
    ws.append(header)

    # 写入数据行
    for ts, direction, byte_list in data_frames:
        row = [ts, direction]
        row.extend(f"{b:02X}" for b in byte_list)
        ws.append(row)

    # 优化列宽
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 8
    for col in range(3, max_byte_len + 3):
        ws.column_dimensions[get_column_letter(col)].width = 6

    wb.save(file_path)
